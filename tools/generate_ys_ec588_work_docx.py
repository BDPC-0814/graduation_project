from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "docs" / "YS-EC588板卡工作总结与成果文档.md"
OUT_PATH = ROOT / "docs" / "YS-EC588板卡工作总结与成果文档.docx"


def set_font(run, *, size: float = 10.5, bold: bool = False, name: str = "微软雅黑", color: RGBColor | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def spacing(p, before: float = 0, after: float = 5, line: float = 1.15) -> None:
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_inline(p, text: str, *, size: float = 10.5) -> None:
    tokens = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            r = p.add_run(token[1:-1])
            set_font(r, size=size, name="Consolas", color=RGBColor(15, 118, 110))
        elif token.startswith("**") and token.endswith("**"):
            r = p.add_run(token[2:-2])
            set_font(r, size=size, bold=True)
        else:
            r = p.add_run(token)
            set_font(r, size=size)


def add_code(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.45)
    spacing(p, before=2, after=7, line=1.0)
    for i, line in enumerate(lines):
        if i:
            p.add_run("\n")
        r = p.add_run(line)
        set_font(r, size=9, name="Consolas", color=RGBColor(30, 41, 59))
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F1F5F9")
    p_pr.append(shd)


def add_md_table(doc: Document, block: list[str]) -> None:
    rows: list[list[str]] = []
    for line in block:
        parts = [p.strip() for p in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", p or "") for p in parts):
            continue
        rows.append(parts)
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            text = row[c_idx] if c_idx < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            spacing(p, after=0, line=1.05)
            if r_idx == 0:
                shade_cell(cell, "0F766E")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text)
                set_font(run, size=9.5, bold=True, color=RGBColor(255, 255, 255))
            else:
                if r_idx % 2 == 0:
                    shade_cell(cell, "F8FAFC")
                run = p.add_run(text)
                set_font(run, size=9)
    doc.add_paragraph()


def setup() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    styles = doc.styles
    styles["Normal"].font.name = "微软雅黑"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        styles[name].font.name = "微软雅黑"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return doc


def convert(doc: Document, md: str) -> None:
    lines = md.splitlines()
    i = 0
    code = False
    code_lines: list[str] = []
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            if code:
                add_code(doc, code_lines)
                code_lines = []
                code = False
            else:
                code = True
            i += 1
            continue
        if code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.strip() == "---":
            p = doc.add_paragraph()
            spacing(p, before=3, after=3)
            r = p.add_run("—" * 38)
            set_font(r, size=8, color=RGBColor(203, 213, 225))
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_md_table(doc, block)
            continue
        m = re.match(r"^(#{1,6})\s+(.+)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                spacing(p, before=8, after=14)
                r = p.add_run(text)
                set_font(r, size=20, bold=True, color=RGBColor(31, 41, 55))
            else:
                p = doc.add_paragraph(style=f"Heading {min(level, 4)}")
                spacing(p, before=8 if level == 2 else 4, after=4)
                r = p.add_run(text)
                set_font(r, size={2: 15, 3: 12.5, 4: 11}.get(level, 10.5), bold=True, color=RGBColor(15, 118, 110) if level == 2 else RGBColor(31, 41, 55))
            i += 1
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            spacing(p, before=2, after=6)
            add_inline(p, line.lstrip("> ").strip(), size=11)
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(15, 118, 110)
            i += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            spacing(p, after=2, line=1.08)
            add_inline(p, bullet.group(1), size=10.5)
            i += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            spacing(p, after=2, line=1.08)
            add_inline(p, numbered.group(1), size=10.5)
            i += 1
            continue
        p = doc.add_paragraph()
        spacing(p, after=5, line=1.15)
        add_inline(p, line, size=10.5)
        i += 1
    if code_lines:
        add_code(doc, code_lines)


def main() -> None:
    doc = setup()
    convert(doc, MD_PATH.read_text(encoding="utf-8"))
    doc.core_properties.title = "YS-EC588板卡工作总结与成果文档"
    doc.core_properties.author = "Codex"
    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
